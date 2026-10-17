"""Create/update Supabase Auth accounts from the team's initial-password workbook.

The workbook is read locally. Passwords are sent directly to Supabase Auth and
are never written to this repository or printed to the terminal.
"""

import json
import os
import sys
import urllib.error
import urllib.parse
import urllib.request

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit("openpyxl이 필요합니다. 먼저 `py -m pip install openpyxl`을 실행하세요.") from exc


SUPABASE_URL = "https://sxtfqejlvhtnkfxlnsdt.supabase.co"


def api_request(path, key, method="GET", body=None):
    data = json.dumps(body).encode("utf-8") if body is not None else None
    request = urllib.request.Request(
        SUPABASE_URL + path,
        data=data,
        method=method,
        headers={
            "apikey": key,
            "Authorization": "Bearer " + key,
            "Content-Type": "application/json",
        },
    )
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as error:
        # Avoid printing response bodies, since they may include account details.
        raise RuntimeError(f"Supabase 요청 실패 (HTTP {error.code})") from None


def account_email(name):
    hex_name = "".join(format(ord(char), "x") for char in name)
    return f"member-{hex_name}@weekly.local"


def main():
    if len(sys.argv) not in (2, 3) or (len(sys.argv) == 3 and sys.argv[2] != "--reset-passwords"):
        raise SystemExit('사용법: py scripts\\provision_auth_users.py "D:\\경로\\초기비밀번호.xlsx" [--reset-passwords]')
    reset_passwords = len(sys.argv) == 3
    key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "").strip()
    if not key:
        raise SystemExit("SUPABASE_SERVICE_ROLE_KEY 환경 변수에 Supabase service_role 비밀키를 지정하세요. 이 키를 공유하거나 저장소에 올리지 마세요.")

    workbook = load_workbook(sys.argv[1], read_only=True, data_only=True)
    sheet = workbook.active
    credentials = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        name = str(row[1] or "").strip() if len(row) > 1 else ""
        password = str(row[3] or "").strip() if len(row) > 3 else ""
        if name and password:
            credentials.append((name, password))
    workbook.close()
    if not credentials:
        raise SystemExit("엑셀에서 성명과 비밀번호 데이터를 찾지 못했습니다.")

    members = api_request("/rest/v1/members?select=id,name,is_admin,status", key)
    by_name = {member["name"]: member for member in members}
    missing = [name for name, _ in credentials if name not in by_name or by_name[name]["status"] != "ACTIVE"]
    if missing:
        raise SystemExit(f"활성 팀원 명단에서 찾을 수 없는 계정이 있습니다 ({len(missing)}명). Supabase members 명단을 먼저 확인하세요.")

    users_payload = api_request("/auth/v1/admin/users?page=1&per_page=1000", key)
    users = users_payload.get("users", []) if isinstance(users_payload, dict) else users_payload
    users_by_email = {user.get("email", "").lower(): user for user in users}
    created = updated = 0

    for name, password in credentials:
        member = by_name[name]
        email = account_email(name)
        metadata = {
            "member_id": member["id"],
            "is_manager": bool(member.get("is_admin")),
            "work_type": member.get("work_type", ""),
        }
        body = {"email": email, "email_confirm": True, "app_metadata": metadata}
        existing = users_by_email.get(email.lower())
        if existing:
            if reset_passwords:
                body["password"] = password
            api_request("/auth/v1/admin/users/" + urllib.parse.quote(existing["id"], safe=""), key, "PUT", body)
            updated += 1
        else:
            body["password"] = password
            created_user = api_request("/auth/v1/admin/users", key, "POST", body)
            users_by_email[email.lower()] = created_user
            created += 1

    print(f"완료: 초기 비밀번호 계정 {len(credentials)}개 처리 (신규 {created}개, 갱신 {updated}개). 비밀번호 값은 표시하지 않았습니다.")


if __name__ == "__main__":
    try:
        main()
    except (OSError, RuntimeError, ValueError, KeyError) as error:
        raise SystemExit(str(error)) from None
